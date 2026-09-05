"""Artifacts render into valid, non-trivial PDF/PPTX files with real content."""

from __future__ import annotations

from kirana.artifacts.analysis_deck import render_deck
from kirana.artifacts.invoice_pdf import render_invoice
from kirana.artifacts.sales_report_pdf import render_sales_report_pdf
from kirana.db.connection import today_iso
from tests.conftest import add

CHAT = 1


def _finalize_sample(repo):
    m = add(repo, "Maggi 70g", "packet", 18, 14, 12, stock=50, brand="Nestle", hsn="1902")
    a = add(repo, "Aashirvaad Atta 5kg", "packet", 5, 270, 250, stock=6, hsn="1101")
    repo.set_shop_profile(name="Test Kirana", gstin="29ABCDE1234F1Z5", address="MG Road")
    repo.cart_add_item(CHAT, m.id, 3)
    repo.cart_add_item(CHAT, a.id, 1)
    repo.cart_set_payment(CHAT, "upi", "UPI-123")
    return repo.finalize_bill(CHAT, "art-key")


def test_invoice_pdf(repo, tmp_path):
    bill = _finalize_sample(repo)
    path = render_invoice(repo, bill.id, out_dir=str(tmp_path))
    data = open(path, "rb").read()
    assert data[:5] == b"%PDF-"
    assert len(data) > 3000


def test_analysis_deck_has_real_charts(repo, tmp_path):
    _finalize_sample(repo)
    d = today_iso()
    path = render_deck(repo, d, d, out_dir=str(tmp_path), title="Test Deck")
    assert open(path, "rb").read(2) == b"PK"  # pptx is a zip

    from pptx import Presentation

    prs = Presentation(path)
    assert len(prs.slides) >= 6
    charts = sum(1 for s in prs.slides for sh in s.shapes if sh.has_chart)
    assert charts >= 3  # native, editable charts (not images)


def test_deck_handles_empty_period(repo, tmp_path):
    add(repo, "Maggi 70g", "packet", 18, 14, 12, stock=50)
    path = render_deck(repo, "2000-01-01", "2000-01-02", out_dir=str(tmp_path))
    from pptx import Presentation

    assert len(Presentation(path).slides) >= 6  # still a valid deck


def test_multilingual_deck(repo, tmp_path):
    _finalize_sample(repo)
    d = today_iso()
    for lang in ("ta", "hi", "en"):
        path = render_deck(repo, d, d, out_dir=str(tmp_path), language=lang)
        from pptx import Presentation
        prs = Presentation(path)
        assert len(prs.slides) >= 6
        assert open(path, "rb").read(2) == b"PK"


def test_sales_report_pdf(repo, tmp_path):
    _finalize_sample(repo)
    d = today_iso()
    path = render_sales_report_pdf(repo, d, d, out_dir=str(tmp_path))
    data = open(path, "rb").read()
    assert data[:5] == b"%PDF-"
    assert len(data) > 3000
